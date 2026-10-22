"""Portfolio import (plan S7.1, M2-I): read, plan, commit and roll back a client file.

The three services are the contract; everything above them is parsing.

* `plan_import` reads a CSV or XLSX upload, maps its columns, normalizes and validates
  every row, and persists the dry run as one `ImportBatch(status=planned)` with one
  `ImportRow` per data row and its outcome. Nothing in the portfolio moves.
* `commit_import` creates the planned clients, each carrying `import_batch_id` as
  provenance. It runs in a Celery `TenantTask` (`apps/clients/tasks.py`); rows that
  fail to write leave the batch `partial` and can be retried by committing again.
* `rollback_import` deletes what a batch created while nothing depends on it. The
  dependents registry below decides what "depends" means, and later todos extend it
  (`register_rollback_dependent`) — AXR1-009 registers service cycles.

Hard limits, all refusals rather than truncations: 5 MiB per file, 5,000 data rows,
and for XLSX a bound on the decompressed archive before openpyxl touches it (a small
zip can inflate into gigabytes of shared strings). openpyxl runs in read-only mode
with `data_only=True`: formulas are never evaluated, only their cached text is read,
and any text a spreadsheet would execute is defused again in the error report.

Dedupe is by normalized CNPJ inside the importing firm ONLY. Another firm holding the
same CNPJ is invisible here and irrelevant (plan Must NOT: no cross-tenant dedupe, no
name-similarity merging), and a CNPJ already on this firm's books is skipped — an
import never rewrites an existing client and never changes a CNPJ.
"""

import codecs
import csv
import hashlib
import io
import logging
import unicodedata
import zipfile
from collections.abc import Callable, Iterable, Iterator, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any, Final
from uuid import UUID

from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import UploadedFile
from django.db import IntegrityError, transaction
from django.db.models import Count, QuerySet
from django.http import Http404
from django.utils import timezone
from django.utils.translation import gettext as _
from django.utils.translation import gettext_lazy

from apps.accounts.models import User
from apps.audit.models import AuditAction
from apps.audit.services import ObjectRef, record_event
from apps.authz.portfolio import PortfolioScope, portfolio_scope
from apps.authz.services import can
from apps.clients.exports import defuse_formula
from apps.clients.maintenance import open_initial_claim
from apps.clients.models import (
    ClientAssignment,
    ClientCompany,
    ClientContact,
    Engagement,
    ImportBatch,
    ImportFormat,
    ImportRow,
    ImportStatus,
    OnboardingItem,
    OnboardingStatus,
    RowOutcome,
    TaxProfile,
)
from apps.core.tenancy import MissingTenantContext, current_tenant_id
from apps.fiscal.models import UF
from apps.fiscal.municipalities import validate_municipality_code
from apps.fiscal.validators import normalize_document, validate_cnpj, validate_cpf
from apps.obligations.models import Document, MonthlyRevenue, Obligation
from apps.obligations.models.operations import ObligationStatus
from apps.tenants.entitlements import lock_entitlement, require_client_capacity
from apps.tenants.models import Invite, Membership

logger = logging.getLogger(__name__)

MAX_ROWS: Final = 5_000
MAX_BYTES: Final = 5 * 1024 * 1024
# XLSX is a zip. Bounds checked on the central directory BEFORE openpyxl reads a
# byte: the declared sizes bound what zipfile will ever inflate (a member that
# inflates past its declared size fails its CRC), so these caps bound memory.
MAX_XLSX_MEMBERS: Final = 500
MAX_XLSX_UNCOMPRESSED: Final = 64 * 1024 * 1024
MAX_XLSX_RATIO: Final = 200
XLSX_MAGIC: Final = b"PK\x03\x04"
CSV_SNIFF_BYTES: Final = 8192

# S8.1 `clients.import` (owner and operations admin). The firm-wide portfolio is
# required as well: an import creates clients nobody is assigned to, and its
# rollback must see every one of them, so a per-user grant of the capability to an
# assigned-only account would create clients it could neither see nor undo.
IMPORT_CAPABILITY: Final = "clients.import"
# Provenance of the initial timeline claim an import opens for each client; the
# rollback dependents registry ignores these rows, which belong to the import.
IMPORT_CLAIM_PROVENANCE: Final = "import"
IMPORT_PROVENANCE_OBJECT: Final = "clients.ImportBatch"

ROW_LIMIT_MESSAGE: Final = gettext_lazy(
    "O arquivo tem mais linhas do que o permitido: limite de 5.000 linhas por "
    "importação. Divida o arquivo e importe cada parte separadamente.",
)


@dataclass(frozen=True)
class ImportField:
    """One ClientCompany field a file column can be mapped to."""

    key: str
    label: Any
    required: bool
    aliases: tuple[str, ...]


FIELDS: Final[tuple[ImportField, ...]] = (
    ImportField("cnpj", gettext_lazy("CNPJ"), required=True, aliases=("cnpj",)),
    ImportField(
        "legal_name",
        gettext_lazy("Razão social"),
        required=True,
        aliases=("razao social", "nome empresarial", "nome", "legal name", "empresa"),
    ),
    ImportField(
        "trade_name",
        gettext_lazy("Nome fantasia"),
        required=False,
        aliases=("nome fantasia", "fantasia", "trade name"),
    ),
    ImportField(
        "cpf",
        gettext_lazy("CPF do titular"),
        required=False,
        aliases=("cpf", "cpf titular"),
    ),
    ImportField(
        "state", gettext_lazy("UF"), required=False, aliases=("uf", "estado", "state")
    ),
    ImportField(
        "municipality_ibge_code",
        gettext_lazy("Código IBGE do município"),
        required=False,
        aliases=(
            "codigo ibge",
            "ibge",
            "municipio ibge",
            "cod ibge",
            "municipality ibge code",
        ),
    ),
    ImportField(
        "main_cnae",
        gettext_lazy("CNAE principal"),
        required=False,
        aliases=("cnae principal", "cnae", "main cnae"),
    ),
)
FIELD_KEYS: Final = frozenset(field.key for field in FIELDS)
REQUIRED_KEYS: Final = frozenset(field.key for field in FIELDS if field.required)

OUTCOME_LABELS: Final[dict[str, Any]] = {
    RowOutcome.CREATED: gettext_lazy("Será criado"),
    RowOutcome.UPDATED: gettext_lazy("Atualizado"),
    RowOutcome.DUPLICATE_SKIPPED: gettext_lazy("Duplicado (ignorado)"),
    RowOutcome.INVALID: gettext_lazy("Inválido"),
}


# ---------------------------------------------------------------------------
# Errors
# ---------------------------------------------------------------------------


class ImportFileError(Exception):
    """The file cannot be planned at all (HTTP 422); nothing was persisted."""

    def __init__(self, message: object) -> None:
        """Keep the reader-facing pt-BR sentence."""
        super().__init__(str(message))
        self.message = str(message)


class MappingIncompleteError(ImportFileError):
    """A required field has no column; the reader must choose one (HTTP 422)."""

    def __init__(
        self,
        message: object,
        *,
        headers: list[str],
        mapping: dict[str, str],
    ) -> None:
        """Carry the file's headers and the proposed mapping back to the form."""
        super().__init__(message)
        self.headers = headers
        self.mapping = mapping


class ImportConflictError(Exception):
    """The batch moved since it was read, or is in the wrong state (HTTP 409)."""


@dataclass(frozen=True)
class Dependent:
    """Records blocking a rollback: `count` of `label` on one imported client."""

    kind: str
    label: str
    client_id: UUID
    count: int


class RollbackRefusedError(Exception):
    """The batch's clients have dependents; nothing was deleted."""

    def __init__(self, dependents: list[Dependent]) -> None:
        """Keep the blocking list for the page that explains the refusal."""
        super().__init__(f"{len(dependents)} blocking dependent group(s)")
        self.dependents = dependents


# ---------------------------------------------------------------------------
# Dependents registry
# ---------------------------------------------------------------------------

DependentFinder = Callable[[UUID, Sequence[UUID]], QuerySet[Any]]


@dataclass(frozen=True)
class DependentKind:
    """A kind of record that, once it exists on an imported client, blocks rollback.

    `finder(tenant_id, client_ids)` returns a queryset with a `client_id` column.
    It must filter on `tenant_id` itself: some registered tables (invites,
    memberships) are tenancy roots that carry no row-level policy.
    """

    key: str
    label: Any
    finder: DependentFinder


_DEPENDENTS: dict[str, DependentKind] = {}


def register_rollback_dependent(key: str, label: Any, finder: DependentFinder) -> None:  # noqa: ANN401
    """Register a kind of record that blocks rolling an import back.

    Called at import time by the module owning the table (later todos: service
    cycles, requests). A duplicate key is a programming error, not an override.
    """
    if key in _DEPENDENTS:
        msg = f"rollback dependent {key!r} is already registered"
        raise ValueError(msg)
    _DEPENDENTS[key] = DependentKind(key=key, label=label, finder=finder)


def registered_dependents() -> tuple[DependentKind, ...]:
    """Return the registry in registration order."""
    return tuple(_DEPENDENTS.values())


def find_dependents(tenant_id: UUID, client_ids: Sequence[UUID]) -> list[Dependent]:
    """List every registered dependent on any of `client_ids`, grouped per client."""
    if not client_ids:
        return []
    found: list[Dependent] = []
    for kind in _DEPENDENTS.values():
        counts = (
            kind.finder(tenant_id, client_ids)
            .order_by()
            .values("client_id")
            .annotate(n=Count("pk"))
        )
        found.extend(
            Dependent(
                kind=kind.key,
                label=str(kind.label),
                client_id=row["client_id"],
                count=row["n"],
            )
            for row in counts
        )
    return sorted(found, key=lambda dep: (str(dep.client_id), dep.kind))


def _registered_by_this_module() -> None:
    register_rollback_dependent(
        "clients.assignment",
        gettext_lazy("responsáveis atribuídos"),
        lambda tenant_id, ids: ClientAssignment.objects.filter(
            tenant_id=tenant_id,
            client_id__in=ids,
        ),
    )
    register_rollback_dependent(
        "clients.contact",
        gettext_lazy("contatos cadastrados"),
        lambda tenant_id, ids: ClientContact.objects.filter(
            tenant_id=tenant_id,
            client_id__in=ids,
        ),
    )
    register_rollback_dependent(
        "clients.engagement",
        gettext_lazy("serviços contratados"),
        lambda tenant_id, ids: Engagement.objects.filter(
            tenant_id=tenant_id,
            client_id__in=ids,
        ),
    )
    register_rollback_dependent(
        "clients.tax_profile",
        gettext_lazy("registros de perfil tributário"),
        lambda tenant_id, ids: TaxProfile.objects.filter(
            tenant_id=tenant_id,
            client_id__in=ids,
        ).exclude(provenance=IMPORT_CLAIM_PROVENANCE),
    )
    register_rollback_dependent(
        "clients.onboarding_progress",
        gettext_lazy("itens de checklist já trabalhados"),
        lambda tenant_id, ids: OnboardingItem.objects.filter(
            tenant_id=tenant_id,
            client_id__in=ids,
        ).exclude(status=OnboardingStatus.PENDING),
    )
    # tenants_invite and tenants_membership are tenancy roots without row-level
    # policies, so both finders filter on the tenant explicitly.
    register_rollback_dependent(
        "tenants.portal_invite",
        gettext_lazy("convites ao portal"),
        lambda tenant_id, ids: (
            # PLATFORM_QUERY_OK: tenancy root; filtered on tenant_id explicitly.
            Invite.objects.filter(
                tenant_id=tenant_id,
                client_id__in=ids,
            )
        ),
    )
    register_rollback_dependent(
        "tenants.portal_membership",
        gettext_lazy("acessos ao portal"),
        lambda tenant_id, ids: (
            # PLATFORM_QUERY_OK: tenancy root; filtered on tenant_id explicitly.
            Membership.objects.filter(
                tenant_id=tenant_id,
                client_id__in=ids,
            )
        ),
    )
    register_rollback_dependent(
        "obligations.document",
        gettext_lazy("documentos"),
        lambda tenant_id, ids: Document.objects.filter(
            tenant_id=tenant_id,
            client_id__in=ids,
        ),
    )
    register_rollback_dependent(
        "obligations.revenue",
        gettext_lazy("faturamentos mensais"),
        lambda tenant_id, ids: MonthlyRevenue.objects.filter(
            tenant_id=tenant_id,
            client_id__in=ids,
        ),
    )
    # Generated calendar rows are not work and cascade away with the client; an
    # obligation somebody settled (paid, waived) is.
    register_rollback_dependent(
        "obligations.settled",
        gettext_lazy("obrigações baixadas"),
        lambda tenant_id, ids: Obligation.objects.filter(
            tenant_id=tenant_id,
            client_id__in=ids,
            status__in=(ObligationStatus.PAID, ObligationStatus.WAIVED),
        ),
    )


_registered_by_this_module()


# ---------------------------------------------------------------------------
# Reading the file
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class ParsedFile:
    """A file reduced to a header and data rows, each tagged with its row number."""

    filename: str
    file_format: str
    sha256: str
    headers: list[str]
    rows: list[tuple[int, list[object]]]


def _read_bytes(upload: UploadedFile) -> bytes:
    if upload.size is not None and upload.size > MAX_BYTES:
        raise ImportFileError(_("O arquivo passa do limite de 5 MiB."))
    data = bytes(upload.read(MAX_BYTES + 1))
    if len(data) > MAX_BYTES:
        raise ImportFileError(_("O arquivo passa do limite de 5 MiB."))
    if not data.strip():
        raise ImportFileError(_("O arquivo está vazio."))
    return data


def _decode(data: bytes) -> str:
    """Decode UTF-8 (with or without BOM), UTF-16 with BOM, or Windows-1252.

    Windows-1252 is what Excel writes when a Brazilian office saves "CSV"; five of
    its byte values are undefined, so a binary or mis-declared file still fails.
    """
    if data.startswith(codecs.BOM_UTF8):
        candidates: tuple[tuple[bytes, str], ...] = ((data[3:], "utf-8"),)
    elif data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        candidates = ((data, "utf-16"),)
    else:
        candidates = ((data, "utf-8"), (data, "cp1252"))
    for payload, encoding in candidates:
        try:
            text = payload.decode(encoding)
        except UnicodeDecodeError:
            continue
        if "\x00" in text:
            break
        return text
    raise ImportFileError(
        _(
            "Não foi possível ler o texto do arquivo. Salve-o como CSV UTF-8 ou "
            "como planilha XLSX e envie de novo.",
        ),
    )


def _is_blank(values: Sequence[object]) -> bool:
    return all(value is None or not str(value).strip() for value in values)


def _collect(
    records: Iterable[Sequence[object]],
) -> tuple[list[str], list[tuple[int, list[object]]]]:
    """Split records into header + numbered non-blank rows, enforcing the row cap.

    Stops reading at row 5,001: the rest of an oversized file is never parsed.
    """
    iterator: Iterator[Sequence[object]] = iter(records)
    headers: list[str] | None = None
    rows: list[tuple[int, list[object]]] = []
    for number, record in enumerate(iterator, start=1):
        if headers is None:
            if _is_blank(record):
                continue
            headers = [str(cell).strip() if cell is not None else "" for cell in record]
            continue
        if _is_blank(record):
            continue
        if len(rows) >= MAX_ROWS:
            raise ImportFileError(ROW_LIMIT_MESSAGE)
        rows.append((number, list(record)))
    if headers is None:
        raise ImportFileError(_("O arquivo não tem cabeçalho nem linhas."))
    if not rows:
        raise ImportFileError(_("O arquivo não tem nenhuma linha de dados."))
    return headers, rows


def _read_csv(data: bytes) -> tuple[list[str], list[tuple[int, list[object]]]]:
    text = _decode(data)
    sample = text[:CSV_SNIFF_BYTES]
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
            sample,
            delimiters=";,\t",
        )
    except csv.Error:
        dialect = csv.excel
    try:
        return _collect(csv.reader(io.StringIO(text, newline=""), dialect))
    except csv.Error as error:
        raise ImportFileError(
            _("O CSV está malformado (%(detail)s).") % {"detail": error},
        ) from error


def _check_zip_bounds(data: bytes) -> None:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as error:
        raise ImportFileError(_("O arquivo não é uma planilha XLSX válida.")) from error
    with archive:
        members = archive.infolist()
        if len(members) > MAX_XLSX_MEMBERS:
            raise ImportFileError(_("A planilha tem partes internas demais."))
        total = 0
        for member in members:
            total += member.file_size
            ratio = member.file_size / max(member.compress_size, 1)
            if total > MAX_XLSX_UNCOMPRESSED or (
                ratio > MAX_XLSX_RATIO and member.file_size > MAX_BYTES
            ):
                raise ImportFileError(
                    _(
                        "A planilha descompactada é grande demais para ser "
                        "importada com segurança.",
                    ),
                )


def _read_xlsx(data: bytes) -> tuple[list[str], list[tuple[int, list[object]]]]:
    # Imported here: openpyxl is only needed on this path and is not light.
    from openpyxl import load_workbook  # noqa: PLC0415

    _check_zip_bounds(data)
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except ImportFileError:
        raise
    except Exception as error:
        # openpyxl surfaces a malformed package as any of KeyError, ValueError,
        # TypeError, zipfile and XML parser errors; all mean "not a workbook".
        logger.info("xlsx rejected: %s", type(error).__name__)
        raise ImportFileError(_("O arquivo não é uma planilha XLSX válida.")) from error
    try:
        if not workbook.worksheets:
            raise ImportFileError(_("A planilha não tem nenhuma aba."))
        sheet = workbook.worksheets[0]
        try:
            return _collect(sheet.iter_rows(values_only=True))
        except ImportFileError:
            raise
        except Exception as error:
            logger.info("xlsx rows rejected: %s", type(error).__name__)
            raise ImportFileError(
                _("O arquivo não é uma planilha XLSX válida."),
            ) from error
    finally:
        workbook.close()


def read_upload(upload: UploadedFile) -> ParsedFile:
    """Read an uploaded CSV or XLSX into a header and numbered rows, or refuse."""
    data = _read_bytes(upload)
    name = (upload.name or "").strip()
    lowered = name.lower()
    if lowered.endswith(".xls"):
        raise ImportFileError(
            _("Planilhas .xls antigas não são aceitas. Salve como XLSX ou CSV."),
        )
    if data.startswith(XLSX_MAGIC):
        file_format = ImportFormat.XLSX
        headers, rows = _read_xlsx(data)
    elif lowered.endswith(".xlsx"):
        raise ImportFileError(_("O arquivo não é uma planilha XLSX válida."))
    else:
        file_format = ImportFormat.CSV
        headers, rows = _read_csv(data)
    return ParsedFile(
        filename=name[:255] or "importacao",
        file_format=str(file_format),
        sha256=hashlib.sha256(data).hexdigest(),
        headers=headers,
        rows=rows,
    )


# ---------------------------------------------------------------------------
# Mapping and row normalization
# ---------------------------------------------------------------------------


def _slug(header: str) -> str:
    decomposed = unicodedata.normalize("NFKD", header)
    bare = "".join(char for char in decomposed if not unicodedata.combining(char))
    words = "".join(char if char.isalnum() else " " for char in bare.lower())
    return " ".join(words.split())


def propose_mapping(headers: Sequence[str]) -> dict[str, str]:
    """Match file headers to fields by exact alias, never by fuzzy similarity."""
    by_slug: dict[str, str] = {}
    for header in headers:
        by_slug.setdefault(_slug(header), header)
    proposal: dict[str, str] = {}
    for field in FIELDS:
        for alias in field.aliases:
            if alias in by_slug and by_slug[alias] not in proposal.values():
                proposal[field.key] = by_slug[alias]
                break
    return proposal


def resolve_mapping(
    headers: list[str],
    mapping: dict[str, str] | None,
) -> dict[str, int]:
    """Validate a field->header mapping against the file; return field->column index.

    With no mapping the proposal is used. An incomplete one raises
    `MappingIncompleteError` carrying the headers so the form can ask.
    """
    chosen = {
        key: header
        for key, header in (mapping or propose_mapping(headers)).items()
        if key in FIELD_KEYS and header
    }
    missing = REQUIRED_KEYS - chosen.keys()
    unknown = [header for header in chosen.values() if header not in headers]
    repeated = len(set(chosen.values())) != len(chosen)
    if missing or unknown or repeated:
        labels = ", ".join(str(field.label) for field in FIELDS if field.key in missing)
        message = (
            _("Escolha a coluna de: %(fields)s.") % {"fields": labels}
            if missing
            else _("Cada coluna do arquivo só pode ser usada uma vez.")
        )
        raise MappingIncompleteError(message, headers=headers, mapping=chosen)
    return {key: headers.index(header) for key, header in chosen.items()}


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "sim" if value else "não"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _document(value: object, width: int) -> str:
    """Normalize a CNPJ/CPF cell; restore leading zeros a spreadsheet dropped.

    A numeric cell cannot have carried a leading zero, so left-padding an integer
    to the document width restores exactly what the spreadsheet stripped; the check
    digits then validate or refuse the result like any other input.
    """
    if isinstance(value, bool):
        return normalize_document(_text(value))
    if isinstance(value, int) or (isinstance(value, float) and value.is_integer()):
        return str(int(value)).zfill(width)
    return normalize_document(_text(value))


VALID_UFS: Final = frozenset(str(choice) for choice in UF.values)
TEXT_LIMITS: Final[dict[str, int]] = {"legal_name": 255, "trade_name": 255}


# The fiscal validators speak English to maintainers (their msgids have no pt-BR
# catalogue); the reader of an import report or a form gets these, chosen by code.
_DOCUMENT_MESSAGES: Final[dict[str, Any]] = {
    "length": gettext_lazy("O %(doc)s deve ter %(width)s caracteres."),
    "characters": gettext_lazy("O %(doc)s tem caracteres que não são permitidos."),
    "degenerate": gettext_lazy("O %(doc)s não pode ter todos os caracteres iguais."),
    "check_digit_characters": gettext_lazy(
        "Os dois últimos caracteres do %(doc)s devem ser números.",
    ),
    "check_digits": gettext_lazy("Os dígitos verificadores do %(doc)s não conferem."),
}


def document_messages(error: ValidationError, doc: str) -> list[str]:
    """Return pt-BR sentences for a CNPJ/CPF validation error, by its code."""
    width = 14 if doc == "CNPJ" else 11
    sentences: list[str] = []
    for item in error.error_list:
        code = (item.code or "").removeprefix(doc.lower() + "_")
        template = _DOCUMENT_MESSAGES.get(code)
        sentences.append(
            str(template) % {"doc": doc, "width": width}
            if template is not None
            else _("%(doc)s inválido.") % {"doc": doc},
        )
    return sentences


def _validate_row(  # noqa: C901 - one flat branch per mapped field, in field order
    cells: list[object],
    columns: dict[str, int],
) -> tuple[str, dict[str, str], list[str]]:
    """Normalize one row per M2-I step 3 and validate it per step 4."""

    def cell(key: str) -> object:
        index = columns.get(key)
        if index is None or index >= len(cells):
            return None
        return cells[index]

    errors: list[str] = []
    cnpj = _document(cell("cnpj"), 14)
    if not cnpj:
        errors.append(_("CNPJ ausente."))
    else:
        try:
            validate_cnpj(cnpj)
        except ValidationError as error:
            errors.extend(document_messages(error, "CNPJ"))
    values: dict[str, str] = {}
    for key in ("legal_name", "trade_name"):
        text = " ".join(_text(cell(key)).split())
        if len(text) > TEXT_LIMITS[key]:
            errors.append(
                _("%(field)s passa de %(limit)s caracteres.")
                % {"field": key, "limit": TEXT_LIMITS[key]},
            )
        values[key] = text[: TEXT_LIMITS[key]]
    if not values["legal_name"]:
        errors.append(_("Razão social ausente."))
    cpf = _document(cell("cpf"), 11)
    if cpf:
        try:
            validate_cpf(cpf)
        except ValidationError as error:
            errors.extend(document_messages(error, "CPF"))
    values["cpf"] = cpf
    state = _text(cell("state")).upper()
    if state and state not in VALID_UFS:
        errors.append(_("UF desconhecida: %(uf)s.") % {"uf": state[:8]})
    values["state"] = state[:2]
    ibge = "".join(
        char for char in _text(cell("municipality_ibge_code")) if char.isalnum()
    )
    errors.extend(_municipality_code_errors(ibge))
    values["municipality_ibge_code"] = ibge[:7]
    cnae = "".join(char for char in _text(cell("main_cnae")) if char.isalnum())
    if cnae and not (len(cnae) == 7 and cnae.isdigit()):  # noqa: PLR2004 - CNAE width
        errors.append(_("O CNAE principal deve ter 7 dígitos."))
    values["main_cnae"] = cnae[:7]
    return cnpj[:32], values, errors


# ---------------------------------------------------------------------------
# Services
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class ImportPlan:
    """The dry run's summary; the rows live on the persisted batch."""

    batch_id: UUID
    version: int
    total: int
    created: int
    invalid: int
    duplicate: int


@dataclass(frozen=True)
class ImportResult:
    """What a commit wrote."""

    batch_id: UUID
    status: str
    created: int
    failed: int
    skipped: int


@dataclass(frozen=True)
class RollbackResult:
    """What a rollback deleted."""

    batch_id: UUID
    deleted: int


def _tenant() -> UUID:
    tenant_id = current_tenant_id.get()
    if tenant_id is None:
        msg = "client import runs inside one tenant's context"
        raise MissingTenantContext(msg)
    return tenant_id


def may_import(actor: User, tenant_id: UUID) -> bool:
    """Report whether the actor may import into this firm (S8.1 `clients.import`)."""
    return can(actor, IMPORT_CAPABILITY, tenant_id=tenant_id) and (
        portfolio_scope(actor, tenant_id=tenant_id) is PortfolioScope.ALL
    )


def _require_import(actor: User, tenant_id: UUID) -> None:
    if not may_import(actor, tenant_id):
        msg = f"{actor} may not import clients into this firm."
        raise PermissionDenied(msg)


def plan_import(
    actor: User,
    file: UploadedFile,
    mapping: dict[str, str] | None,
) -> ImportPlan:
    """Dry-run a file: persist one planned batch with one outcome per row.

    Refusals that concern the whole file (size, row cap, encoding, format, missing
    mapping) raise before anything is written.
    """
    tenant_id = _tenant()
    _require_import(actor, tenant_id)
    parsed = read_upload(file)
    columns = resolve_mapping(parsed.headers, mapping)
    planned = [
        (number, *_validate_row(cells, columns)) for number, cells in parsed.rows
    ]
    valid_cnpjs = {cnpj for _n, cnpj, _v, errors in planned if not errors}
    existing = set(
        ClientCompany.objects.filter(
            tenant_id=tenant_id,
            cnpj__in=valid_cnpjs,
        ).values_list("cnpj", flat=True),
    )
    seen: dict[str, int] = {}
    rows: list[ImportRow] = []
    with transaction.atomic():
        batch = ImportBatch.objects.create(
            tenant_id=tenant_id,
            created_by=actor,
            filename=parsed.filename,
            file_format=parsed.file_format,
            file_sha256=parsed.sha256,
            mapping={key: parsed.headers[index] for key, index in columns.items()},
            total_rows=len(planned),
        )
        for number, cnpj, values, found in planned:
            errors = found
            outcome = RowOutcome.CREATED
            if errors:
                outcome = RowOutcome.INVALID
            elif cnpj in existing:
                outcome = RowOutcome.DUPLICATE_SKIPPED
                errors = [_("CNPJ já cadastrado neste escritório.")]
            elif cnpj in seen:
                outcome = RowOutcome.DUPLICATE_SKIPPED
                errors = [
                    _("CNPJ repetido: já aparece na linha %(line)s.")
                    % {"line": seen[cnpj]},
                ]
            else:
                seen[cnpj] = number
            rows.append(
                ImportRow(
                    tenant_id=tenant_id,
                    batch=batch,
                    row_number=number,
                    outcome=outcome,
                    cnpj=cnpj,
                    values=values,
                    errors=errors,
                ),
            )
        ImportRow.objects.bulk_create(rows, batch_size=1000)
        counts = dict.fromkeys(RowOutcome.values, 0)
        for row in rows:
            counts[row.outcome] += 1
        batch.planned_created = counts[RowOutcome.CREATED]
        batch.planned_invalid = counts[RowOutcome.INVALID]
        batch.planned_duplicate = counts[RowOutcome.DUPLICATE_SKIPPED]
        batch.save(
            update_fields=["planned_created", "planned_invalid", "planned_duplicate"],
        )
        record_event(
            action=AuditAction.IMPORT_PLANNED,
            tenant_id=tenant_id,
            actor=actor,
            obj=ObjectRef(type=IMPORT_PROVENANCE_OBJECT, id=str(batch.pk)),
            metadata={
                "format": parsed.file_format,
                "rows": len(rows),
                "created": batch.planned_created,
                "invalid": batch.planned_invalid,
                "duplicate": batch.planned_duplicate,
            },
        )
    return ImportPlan(
        batch_id=batch.pk,
        version=batch.version,
        total=len(rows),
        created=batch.planned_created,
        invalid=batch.planned_invalid,
        duplicate=batch.planned_duplicate,
    )


def _locked_batch(
    actor: User,
    tenant_id: UUID,
    batch_id: UUID,
    expected_version: int,
    allowed: Iterable[str],
) -> ImportBatch:
    batch = ImportBatch.objects.select_for_update().filter(pk=batch_id).first()
    if batch is None:
        raise Http404
    _require_import(actor, tenant_id)
    if batch.version != expected_version:
        msg = f"batch {batch_id} is at version {batch.version}, not {expected_version}"
        raise ImportConflictError(msg)
    if batch.status not in set(allowed):
        msg = f"batch {batch_id} is {batch.status}"
        raise ImportConflictError(msg)
    return batch


def queue_commit(actor: User, batch_id: UUID, *, expected_version: int) -> ImportBatch:
    """Mark a planned (or partially committed) batch as queued for the commit task."""
    tenant_id = _tenant()
    with transaction.atomic():
        tenant = lock_entitlement(tenant_id)
        batch = _locked_batch(
            actor,
            tenant_id,
            batch_id,
            expected_version,
            (ImportStatus.PLANNED, ImportStatus.PARTIAL),
        )
        require_client_capacity(tenant, additional=pending_client_count(batch))
        previous = str(batch.status)
        batch.status = ImportStatus.QUEUED
        batch.version += 1
        batch.save(update_fields=["status", "version"])
        record_event(
            action=AuditAction.IMPORT_QUEUED,
            tenant_id=tenant_id,
            actor=actor,
            obj=ObjectRef(type=IMPORT_PROVENANCE_OBJECT, id=str(batch.pk)),
            metadata={"from": previous, "version": batch.version},
        )
        return batch


def pending_client_count(batch: ImportBatch) -> int:
    """Only new identities need slots; existing or already committed rows do not."""
    return (
        batch.rows.filter(outcome=RowOutcome.CREATED, client__isnull=True)
        .exclude(cnpj__in=ClientCompany.objects.values("cnpj"))
        .count()
    )


def release_failed_commit(
    actor: User,
    batch_id: UUID,
    *,
    expected_version: int,
    error: str,
) -> ImportBatch | None:
    """Hand a queued batch whose commit crashed back to the person importing.

    Called by the commit task after `commit_import` raised something other than a
    conflict; that commit's savepoint already undid every row it wrote. Under the
    batch lock, and only while the batch is still the QUEUED version the task was
    given (a redelivered or late task finds a newer version and changes nothing),
    the batch returns to `partial` when earlier rows were already committed and to
    `planned` otherwise, its version advances, and one event records the failure.

    No capability re-check, deliberately: this grants nothing and undoes nothing a
    person did; it only stops a batch from sitting in a queue that will never drain,
    and it must succeed even if the requester lost the import capability meanwhile.
    The event is attributed to the requester whose commit failed.
    """
    tenant_id = _tenant()
    with transaction.atomic():
        batch = ImportBatch.objects.select_for_update().filter(pk=batch_id).first()
        if (
            batch is None
            or batch.status != ImportStatus.QUEUED
            or batch.version != expected_version
        ):
            return None
        committed = batch.rows.filter(client__isnull=False).exists()
        batch.status = ImportStatus.PARTIAL if committed else ImportStatus.PLANNED
        batch.version += 1
        batch.save(update_fields=["status", "version"])
        record_event(
            action=AuditAction.IMPORT_COMMIT_FAILED,
            tenant_id=tenant_id,
            actor=actor,
            obj=ObjectRef(type=IMPORT_PROVENANCE_OBJECT, id=str(batch.pk)),
            metadata={
                "from": str(ImportStatus.QUEUED),
                "to": str(batch.status),
                "version": batch.version,
                "error": error,
            },
        )
        return batch


def _municipality_code_errors(ibge: str) -> list[str]:
    if ibge and not (len(ibge) == 7 and ibge.isdigit()):  # noqa: PLR2004 - IBGE width
        return [_("O código IBGE do município deve ter 7 dígitos.")]
    try:
        validate_municipality_code(ibge)
    except ValidationError as error:
        return list(error.messages)
    return []


def _create_client(batch: ImportBatch, row: ImportRow) -> ClientCompany:
    values = row.values
    validate_municipality_code(values.get("municipality_ibge_code", ""))
    client = ClientCompany(
        tenant_id=batch.tenant_id,
        cnpj=row.cnpj,
        legal_name=values.get("legal_name", ""),
        trade_name=values.get("trade_name", ""),
        cpf=values.get("cpf") or None,
        state=values.get("state", ""),
        municipality_ibge_code=values.get("municipality_ibge_code", ""),
        main_cnae=values.get("main_cnae", ""),
        import_batch=batch,
    )
    client.full_clean(exclude=["tenant", "import_batch"], validate_unique=False)
    with transaction.atomic():
        client.save()
        # The flag and the timeline agree from the first instant (T6 invariant);
        # the claim is attributed to whoever planned the import.
        open_initial_claim(
            client=client,
            recorded_by=batch.created_by,
            provenance=IMPORT_CLAIM_PROVENANCE,
        )
    return client


def commit_import(
    actor: User, batch_id: UUID, *, expected_version: int
) -> ImportResult:
    """Create the planned clients, each stamped with the batch as provenance.

    Runs under the batch row lock with the capability re-checked inside it. Each row
    writes in its own savepoint: a CNPJ that reached the firm after the dry run is
    skipped as a duplicate, and any other write failure marks the row for retry and
    leaves the batch `partial` instead of undoing the rows that succeeded.
    """
    tenant_id = _tenant()
    with transaction.atomic():
        tenant = lock_entitlement(tenant_id)
        batch = _locked_batch(
            actor,
            tenant_id,
            batch_id,
            expected_version,
            (ImportStatus.PLANNED, ImportStatus.QUEUED, ImportStatus.PARTIAL),
        )
        pending = list(
            batch.rows.select_for_update()
            .filter(outcome=RowOutcome.CREATED, client__isnull=True)
            .order_by("row_number"),
        )
        existing = set(
            ClientCompany.objects.filter(
                tenant_id=tenant_id,
                cnpj__in=[row.cnpj for row in pending],
            ).values_list("cnpj", flat=True),
        )
        require_client_capacity(
            tenant, additional=sum(row.cnpj not in existing for row in pending)
        )
        created = failed = skipped = 0
        for row in pending:
            if row.cnpj in existing:
                row.outcome = RowOutcome.DUPLICATE_SKIPPED
                row.errors = [_("CNPJ cadastrado no escritório depois da prévia.")]
                row.commit_error = ""
                row.save(update_fields=["outcome", "errors", "commit_error"])
                skipped += 1
                continue
            try:
                client = _create_client(batch, row)
            except (IntegrityError, ValidationError) as error:
                logger.warning(
                    "import row failed",
                    extra={"batch_id": str(batch.pk), "row": row.row_number},
                )
                row.commit_error = _("Falha ao gravar: %(kind)s") % {
                    "kind": type(error).__name__,
                }
                row.save(update_fields=["commit_error"])
                failed += 1
                continue
            row.client = client
            row.commit_error = ""
            row.save(update_fields=["client", "commit_error"])
            created += 1
        batch.committed_count = batch.rows.filter(client__isnull=False).count()
        batch.failed_count = failed
        batch.planned_duplicate += skipped
        batch.planned_created -= skipped
        batch.status = ImportStatus.PARTIAL if failed else ImportStatus.COMMITTED
        batch.committed_at = timezone.now()
        batch.version += 1
        batch.save(
            update_fields=[
                "committed_count",
                "failed_count",
                "planned_duplicate",
                "planned_created",
                "status",
                "committed_at",
                "version",
            ],
        )
        record_event(
            action=AuditAction.IMPORT_COMMITTED,
            tenant_id=tenant_id,
            actor=actor,
            obj=ObjectRef(type=IMPORT_PROVENANCE_OBJECT, id=str(batch.pk)),
            metadata={"created": created, "failed": failed, "skipped": skipped},
        )
        return ImportResult(
            batch_id=batch.pk,
            status=batch.status,
            created=created,
            failed=failed,
            skipped=skipped,
        )


def rollback_import(
    actor: User,
    batch_id: UUID,
    *,
    expected_version: int,
) -> RollbackResult:
    """Delete the clients a batch created, unless anything depends on them.

    The batch AND its clients are locked before the dependents are counted, so a
    write path that locks the client (every `apps.clients.maintenance` service does)
    cannot add a dependent between the count and the delete. A refusal raises
    `RollbackRefusedError` listing the blocking records and writes nothing.
    """
    tenant_id = _tenant()
    with transaction.atomic():
        batch = _locked_batch(
            actor,
            tenant_id,
            batch_id,
            expected_version,
            (ImportStatus.COMMITTED, ImportStatus.PARTIAL),
        )
        client_ids = list(
            ClientCompany.objects.select_for_update()
            .filter(tenant_id=tenant_id, import_batch=batch)
            .values_list("pk", flat=True),
        )
        dependents = find_dependents(tenant_id, client_ids)
        if dependents:
            raise RollbackRefusedError(dependents)
        deleted = len(client_ids)
        ClientCompany.objects.filter(pk__in=client_ids).delete()
        batch.status = ImportStatus.ROLLED_BACK
        batch.rolled_back_at = timezone.now()
        batch.version += 1
        batch.save(update_fields=["status", "rolled_back_at", "version"])
        record_event(
            action=AuditAction.IMPORT_ROLLED_BACK,
            tenant_id=tenant_id,
            actor=actor,
            obj=ObjectRef(type=IMPORT_PROVENANCE_OBJECT, id=str(batch.pk)),
            metadata={"deleted": deleted},
        )
        return RollbackResult(batch_id=batch.pk, deleted=deleted)


# ---------------------------------------------------------------------------
# Error report
# ---------------------------------------------------------------------------

ERROR_COLUMNS: Final = ("linha", "cnpj", "razao_social", "resultado", "motivo")


def build_error_csv(batch: ImportBatch) -> str:
    """Render every row that did not become a client, formula-defused.

    The `apps/clients/exports.py` rules, applied to EVERY column: QUOTE_ALL so a
    spreadsheet keeps a CNPJ as text, and `defuse_formula` on each cell. Unlike the
    registry export the CNPJ column is defused too — here it holds what the file
    said, which for an invalid row may be anything, including `=cmd|...`.
    """
    buffer = io.StringIO()
    writer = csv.writer(buffer, quoting=csv.QUOTE_ALL, lineterminator="\n")
    writer.writerow(ERROR_COLUMNS)
    rows = batch.rows.exclude(outcome=RowOutcome.CREATED, commit_error="").order_by(
        "row_number",
    )
    for row in rows:
        reasons = [*row.errors, *([row.commit_error] if row.commit_error else [])]
        writer.writerow(
            [
                defuse_formula(str(row.row_number)),
                defuse_formula(row.cnpj),
                defuse_formula(str(row.values.get("legal_name", ""))),
                defuse_formula(str(OUTCOME_LABELS.get(row.outcome, row.outcome))),
                defuse_formula(" ".join(str(reason) for reason in reasons)),
            ],
        )
    return buffer.getvalue()


__all__ = [
    "FIELDS",
    "MAX_BYTES",
    "MAX_ROWS",
    "OUTCOME_LABELS",
    "Dependent",
    "DependentKind",
    "ImportConflictError",
    "ImportFileError",
    "ImportPlan",
    "ImportResult",
    "MappingIncompleteError",
    "RollbackRefusedError",
    "RollbackResult",
    "build_error_csv",
    "commit_import",
    "document_messages",
    "find_dependents",
    "may_import",
    "plan_import",
    "propose_mapping",
    "queue_commit",
    "read_upload",
    "register_rollback_dependent",
    "registered_dependents",
    "release_failed_commit",
    "resolve_mapping",
    "rollback_import",
]
